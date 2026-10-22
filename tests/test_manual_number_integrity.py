# -*- coding: utf-8 -*-
"""T1 回归：自动编号接管手工前缀（「一、」「1.」）时，文本一致性校验不得误判硬失败。

审计 R-01 的根因链：
- core/style_set.py 会剥掉手工前缀，但只返回 bool，真实剥掉的前缀（含标点/空格）丢失；
- core/agent.py 用 core/extract.py 的 manual_number 标签（"1"）拼白名单，
  中文前缀 extract 根本识别不到，数字前缀则丢了 ". "；
- core/text_integrity.py 按这个不完整前缀从原稿与终稿两边剥：
  原稿 "1. 标题" 变成 ". 标题"，终稿 "标题" 不变，于是被判成删一段/加一段。
"""

import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core.agent import Agent
from core.apply import apply_format
from core.text_integrity import check_text_integrity


def _document(path, *texts):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    doc.save(path)
    return path


def _spec(heading_role, numbering):
    return {"roles": {
        "title": {
            "font_eastasia": "SimHei", "size_pt": 22, "alignment": "center"},
        heading_role: {
            "font_eastasia": "SimHei", "size_pt": 16, "alignment": "left",
            "numbering": numbering,
        },
        "body": {
            "font_eastasia": "SimSun", "size_pt": 12, "alignment": "left"},
    }}


# (角色, 源稿标题文字, 期望输出文字, 实际手工前缀, numbering spec)
CASES = [
    pytest.param(
        "heading_1", "一、第一章 绪论", "第一章 绪论", "一、",
        {"group": "headings", "level": 0, "num_format": "chineseCounting",
         "level_text": "%1、"},
        id="chinese-number-prefix",
    ),
    pytest.param(
        "heading_3", "1. 第一章 绪论", "第一章 绪论", "1. ",
        {"group": "headings", "level": 0, "num_format": "decimal",
         "level_text": "%1."},
        id="arabic-number-prefix",
    ),
]


@pytest.mark.parametrize(
    "heading_role,source_heading,expected_heading,prefix,numbering", CASES)
def test_manual_prefix_replaced_by_automatic_numbering_is_not_a_text_change(
        tmp_path, heading_role, source_heading, expected_heading, prefix, numbering):
    source = _document(
        tmp_path / "source.docx", "文档标题", source_heading, "正文第一段")
    roles = {0: "title", 1: heading_role, 2: "body"}

    result = Agent().run(
        source, tmp_path / "out.docx",
        spec=_spec(heading_role, numbering), rolemap=roles)

    assert result["text_integrity"]["ok"], result["text_integrity"]
    output = Document(result["out_path"])
    texts = [paragraph.text for paragraph in output.paragraphs]
    assert texts == ["文档标题", expected_heading, "正文第一段"]

    # 手工前缀是被自动编号取代，而不是简单删字：标题样式上必须仍有编号定义。
    heading_style = output.styles[output.paragraphs[1].style.name]
    style_ppr = heading_style.element.find(qn("w:pPr"))
    assert style_ppr is not None and style_ppr.find(qn("w:numPr")) is not None


@pytest.mark.parametrize(
    "heading_role,source_heading,expected_heading,prefix,numbering", CASES)
def test_changelog_reports_the_exact_stripped_prefix(
        tmp_path, heading_role, source_heading, expected_heading, prefix, numbering):
    source = _document(
        tmp_path / "source.docx", "文档标题", source_heading, "正文第一段")
    roles = {0: "title", 1: heading_role, 2: "body"}

    changelog = apply_format(
        source, _spec(heading_role, numbering), roles, tmp_path / "out.docx")

    change = next(item for item in changelog if item["idx"] == 1)
    assert "manual_number_prefix_removed" in change["changed_fields"]
    # 必须报告实际剥掉的完整前缀（含标点/空格），而不是 extract 的纯 label。
    assert change["stripped_prefixes"] == [prefix]
    assert Document(tmp_path / "out.docx").paragraphs[1].text == expected_heading


def test_incomplete_or_unknown_prefix_is_not_silently_allowed(tmp_path):
    source = _document(tmp_path / "source.docx", "1. 第一章 绪论")
    output = _document(tmp_path / "out.docx", "第一章 绪论")

    # 没有声明前缀：删字仍是硬失败。
    assert check_text_integrity(source, output)["ok"] is False
    # 只有 label、丢了 ". "：不匹配，仍然硬失败（不放宽成“任意前缀都能剥”）。
    assert check_text_integrity(
        source, output, expected_stripped_prefixes=["1"])["ok"] is False
    # 声明了别的前缀：同样不匹配。
    assert check_text_integrity(
        source, output, expected_stripped_prefixes=["2. "])["ok"] is False
    # 只有完整前缀才允许抵消这段差异。
    assert check_text_integrity(
        source, output, expected_stripped_prefixes=["1. "])["ok"] is True
