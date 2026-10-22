"""R2-2 回归：域刷新改写最终交付物后必须重新通过文本完整性校验。

审计 R-11/B5 的缺口：core/agent.py 原先「先完整性校验 → 再 refresh_fields
（Word 改写目录/页码/页眉）→ 直接提交」，刷新改了字不会被复检；视觉修复
重排（repair_transaction）路径同样是先校验后刷新、未复检。

本文件全程 mock core.field_refresh.refresh_fields_word，绝不拉起 Word/COM；
真实 Word 真机域刷新本轮不实测（人工验收记 not_run）。
"""

import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import core.field_refresh as field_refresh_module
import core.verify_visual as visual_module
from core.agent import Agent
from core.safe_output import IntegrityViolationError


def _spec():
    return {"roles": {"body": {
        "font_eastasia": "SimSun", "size_pt": 12, "alignment": "left",
    }}}


def _document(path, *texts):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    doc.save(path)
    return path


def _corrupt(path):
    """模拟 Word 刷新顺带改写一处正文文字。"""
    doc = Document(path)
    doc.paragraphs[0].text = "Rewritten by field refresh"
    doc.save(path)


def _fail_evidence(tmp_path, refresh_calls):
    """旧实现照常提交时的证据串：交付物是否被发布、刷新调用是否发生。"""
    return (
        "域刷新改字后必须触发完整性硬失败且不发布交付物；"
        "旧实现照常提交："
        f"out_path exists={(tmp_path / 'result.docx').exists()}, "
        f"report exists={(tmp_path / 'result_report.md').exists()}, "
        f"tracked exists={(tmp_path / 'result_tracked.docx').exists()}, "
        f"refresh_calls={list(refresh_calls)}"
    )


def test_main_flow_refresh_text_change_is_hard_failure(tmp_path, monkeypatch):
    """主流程：刷新改字 → IntegrityViolationError，且四件套不发布。"""
    source = _document(tmp_path / "source.docx", "Original body")
    out = tmp_path / "result.docx"
    refresh_calls = []

    def corrupting_refresh(path):
        refresh_calls.append(str(path))
        _corrupt(path)
        return {"total": 1, "updated": 1}

    monkeypatch.setattr(
        field_refresh_module, "refresh_fields_word", corrupting_refresh)

    try:
        Agent(llm=object()).run(
            source, out, spec=_spec(), rolemap={0: "body"},
            refresh_fields=True)
    except IntegrityViolationError:
        assert not out.exists(), "硬失败后不得发布排版稿"
        assert not (tmp_path / "result_report.md").exists()
        assert not (tmp_path / "result_report.docx").exists()
        assert not (tmp_path / "result_tracked.docx").exists()
        assert len(refresh_calls) == 1
    else:
        pytest.fail(_fail_evidence(tmp_path, refresh_calls))


def test_repair_flow_refresh_text_change_is_hard_failure(tmp_path, monkeypatch):
    """视觉修复重排路径：修复后刷新改字 → IntegrityViolationError，不发布。"""
    source = _document(tmp_path / "source.docx", "Original body")
    out = tmp_path / "result.docx"
    refresh_calls = []

    def verify(*args, **kwargs):
        return [{"role": "body", "field": "size_pt", "pass": False,
                 "observed": "12", "expected": "14"}]

    def selective_refresh(path):
        refresh_calls.append(str(path))
        if len(refresh_calls) >= 2:
            # 第 1 次是主流程候选稿刷新（不改字）；第 2 次是修复重排稿刷新，
            # 只让它改字，以覆盖“修复重排→刷新→未复检”的缺口路径。
            _corrupt(path)
        return {"total": 0, "updated": 0}

    monkeypatch.setattr(visual_module, "verify_visual", verify)
    monkeypatch.setattr(
        field_refresh_module, "refresh_fields_word", selective_refresh)

    try:
        Agent(llm=object()).run(
            source, out, spec=_spec(), rolemap={0: "body"},
            verify=True, refresh_fields=True)
    except IntegrityViolationError:
        assert not out.exists(), "硬失败后不得发布修复重排稿"
        assert not (tmp_path / "result_report.md").exists()
        assert not (tmp_path / "result_tracked.docx").exists()
        # 主流程与修复重排的两次刷新都真实执行过（两条路径均被覆盖）
        assert len(refresh_calls) == 2
    else:
        pytest.fail(_fail_evidence(tmp_path, refresh_calls))


def test_normal_refresh_records_post_refresh_integrity(tmp_path, monkeypatch):
    """正常刷新（不改字）：result 记录刷新后复检结论，既有字段不变。"""
    source = _document(tmp_path / "source.docx", "Original body")
    out = tmp_path / "result.docx"

    def no_op_refresh(path):
        return {"total": 0, "updated": 0}

    monkeypatch.setattr(
        field_refresh_module, "refresh_fields_word", no_op_refresh)
    result = Agent(llm=object()).run(
        source, out, spec=_spec(), rolemap={0: "body"}, refresh_fields=True)

    # 新字段：刷新后复检结论
    assert result["post_refresh_integrity"]["ok"] is True
    # 既有字段保持不变
    assert result["text_integrity"]["ok"] is True
    assert result["field_refresh"] == {"total": 0, "updated": 0}
    assert result["verification_status"] == "not_run"
    for key in ("spec", "rolemap", "stylemap", "changelog", "issues",
                "preflight", "out_path", "report_path", "report_docx_path",
                "tracked_path"):
        assert key in result, f"既有 result 字段 {key!r} 缺失"
    for key in ("out_path", "report_path", "report_docx_path",
                "tracked_path"):
        assert Path(result[key]).exists()


def test_without_refresh_no_refresh_called_and_field_is_none(tmp_path,
                                                              monkeypatch):
    """refresh_fields=False：不调用刷新，post_refresh_integrity 为 None。"""
    source = _document(tmp_path / "source.docx", "Original body")
    out = tmp_path / "result.docx"

    def forbidden(path):
        pytest.fail("refresh_fields=False 时不得调用 Word 刷新")

    monkeypatch.setattr(field_refresh_module, "refresh_fields_word", forbidden)
    result = Agent(llm=object()).run(
        source, out, spec=_spec(), rolemap={0: "body"}, refresh_fields=False)

    assert result["post_refresh_integrity"] is None
    assert result["field_refresh"] is None
    assert Path(result["out_path"]).exists()
