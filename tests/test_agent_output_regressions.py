"""Regression coverage for template protection, TOC integrity, and repair rollback."""

import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import core.agent as agent_module
import core.report_docx as report_module
import core.verify_visual as visual_module
from core.agent import Agent
from core.safe_output import IntegrityViolationError, UnsafeOutputPathError


def _spec(toc=False):
    spec = {"roles": {"body": {
        "font_eastasia": "SimSun", "size_pt": 12, "alignment": "left",
    }}}
    if toc:
        spec["toc"] = {"enabled": True}
    return spec


def _document(path, *texts):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    doc.save(path)
    return path


@pytest.mark.parametrize("artifact", ["main", "markdown", "report", "tracked"])
def test_template_cannot_be_overwritten_by_any_artifact(tmp_path, artifact):
    source = _document(tmp_path / "source.docx", "Original body")
    output = tmp_path / "result.docx"
    template = {
        "main": output,
        "markdown": tmp_path / "report.docx",
        "report": tmp_path / "result_report.docx",
        "tracked": tmp_path / "result_tracked.docx",
    }[artifact]
    _document(template, "Reusable template")
    original = template.read_bytes()

    with pytest.raises(UnsafeOutputPathError, match="参考模板"):
        Agent().run(
            source, output, template_path=template,
            template_rolemap={0: "body"}, rolemap={0: "body"},
            report_path=template if artifact == "markdown" else None,
        )

    assert template.read_bytes() == original
    assert not list(tmp_path.glob(".*.tmp*"))


@pytest.mark.parametrize("heading", [False, True])
def test_toc_integrity_uses_actual_insertions(tmp_path, heading):
    source = _document(tmp_path / "source.docx", "目录", "First chapter")
    spec = _spec(toc=True)
    roles = {0: "body", 1: "heading_1" if heading else "body"}
    result = Agent().run(source, tmp_path / "result.docx", spec=spec, rolemap=roles)

    texts = [p.text for p in Document(result["out_path"]).paragraphs]
    assert result["text_integrity"]["ok"]
    assert texts.count("目录") == (2 if heading else 1)
    assert texts[-1] == "First chapter"


def _mock_visual_repair(monkeypatch, tmp_path):
    transactions = []
    original_transaction = agent_module.AtomicOutputSet

    def transaction(paths):
        value = original_transaction(paths)
        transactions.append(value)
        return value

    original_bytes = {}

    def verify(*args, **kwargs):
        original_bytes.update({
            key: Path(path).read_bytes()
            for key, path in transactions[0].temp_paths.items()
        })
        return [{"role": "body", "field": "size_pt", "pass": False,
                 "observed": "12", "expected": "14"}]

    monkeypatch.setattr(agent_module, "AtomicOutputSet", transaction)
    monkeypatch.setattr(visual_module, "verify_visual", verify)
    source = _document(tmp_path / "source.docx", "Original body")
    return source, transactions, original_bytes


@pytest.mark.parametrize("stage", ["main", "markdown", "report", "tracked"])
def test_failed_repair_keeps_entire_validated_output_set(tmp_path, monkeypatch, stage):
    source, transactions, original_bytes = _mock_visual_repair(monkeypatch, tmp_path)
    owner, name, failing_call, path_arg = {
        "main": (agent_module, "apply_format", 3, 3),
        "markdown": (agent_module, "write_report", 2, 2),
        "report": (report_module, "build_report_docx", 2, 2),
        "tracked": (agent_module, "apply_format", 4, 3),
    }[stage]
    original = getattr(owner, name)
    calls = 0

    def fail_save(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == failing_call:
            Path(args[path_arg]).write_bytes(b"partial failed save")
            raise OSError("simulated save failure")
        return original(*args, **kwargs)

    monkeypatch.setattr(owner, name, fail_save)
    result = Agent(llm=object()).run(
        source, tmp_path / "result.docx", spec=_spec(), rolemap={0: "body"}, verify=True,
    )

    for key, expected in original_bytes.items():
        assert Path(result[key]).read_bytes() == expected
    assert result["spec"]["roles"]["body"]["size_pt"] == 12
    assert result["applied_fixes"] == []
    assert result["text_integrity"]["ok"]
    assert not any(Path(p).exists() for p in transactions[1].temp_paths.values())


def test_successful_repair_commits_matching_spec_and_artifacts(tmp_path, monkeypatch):
    source, transactions, _ = _mock_visual_repair(monkeypatch, tmp_path)
    result = Agent(llm=object()).run(
        source, tmp_path / "result.docx", spec=_spec(), rolemap={0: "body"}, verify=True,
    )

    assert result["spec"]["roles"]["body"]["size_pt"] == 14
    assert len(result["applied_fixes"]) == 1
    for key in ("out_path", "tracked_path"):
        doc = Document(result[key])
        assert doc.paragraphs[0].text == "Original body"
        assert doc.paragraphs[0].style.font.size.pt == 14
    assert result["text_integrity"]["ok"]
    assert all(Path(p).exists() for p in transactions[1].final_paths.values())


def test_repair_content_changes_remain_a_hard_failure(tmp_path, monkeypatch):
    source, _, _ = _mock_visual_repair(monkeypatch, tmp_path)
    output = _document(tmp_path / "result.docx", "Existing final")
    before = output.read_bytes()
    original = agent_module.apply_format
    calls = 0

    def change_text(*args, **kwargs):
        nonlocal calls
        calls += 1
        changes = original(*args, **kwargs)
        if calls == 3:
            doc = Document(args[3])
            doc.add_paragraph("Unexpected content")
            doc.save(args[3])
        return changes

    monkeypatch.setattr(agent_module, "apply_format", change_text)
    with pytest.raises(IntegrityViolationError):
        Agent(llm=object()).run(
            source, output, spec=_spec(), rolemap={0: "body"}, verify=True,
        )
    assert output.read_bytes() == before
    assert not (tmp_path / "result_report.md").exists()
