"""T2 回归：track=True 修订稿必须为原本没有直接格式的段落/run 补写修订标记。

审计 R-02 的根因：
- core/track_changes.snapshot_paragraph 对缺失的 w:pPr / w:rPr 返回 None；
- mark_paragraph_revision 看到 None 直接 continue/跳过；
- 于是纯文本段落（无 pPr）与无 rPr 的 run 在修订稿里没有任何
  w:pPrChange / w:rPrChange，Word 里“拒绝全部修订”无法回到原样；
- 另外段落排版会拆分/插入 run（label_prefix），旧快照若按 zip 下标对应，
  会把旧属性错配给新插入的 run，甚至漏掉后面的原 run。

本文件只断言「存在修订标记 + 旧快照完整且正确对应」这些结构性事实。
“Word/WPS 拒绝全部修订后格式回到原样”属人工验收，未执行时记录 not_run。
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core.apply import apply_format

BODY_SPEC = {"roles": {"body": {
    "font_eastasia": "仿宋_GB2312",
    "font_ascii": "Times New Roman",
    "size_pt": 16,
    "alignment": "justify",
    "first_line_indent_chars": 2,
}}}

LABEL_SPEC = {"roles": {"body": {
    "font_eastasia": "仿宋_GB2312",
    "font_ascii": "Times New Roman",
    "size_pt": 16,
    "label_prefix": {"text": ["摘要："], "bold": True},
}}}


def _change(element, tag):
    """取 element 下指定修订记录（不存在返回 None）。"""
    return None if element is None else element.find(qn(tag))


def _old_state(change, tag):
    """取修订记录里的旧状态元素（w:pPr / w:rPr）。"""
    assert change is not None
    state = change.find(qn(tag))
    assert state is not None, "修订记录必须携带完整的旧状态快照"
    return state


def test_plain_paragraph_and_run_get_revision_markers(tmp_path):
    """无 pPr 的纯文本段落、无 rPr 的 run 也要有修订标记与空旧状态。"""
    source = tmp_path / "src.docx"
    tracked = tmp_path / "tracked.docx"

    document = Document()
    document.add_paragraph("完全没有直接格式的正文。")
    document.save(source)

    apply_format(str(source), BODY_SPEC, {0: "body"}, str(tracked), track=True)

    paragraph = Document(tracked).paragraphs[0]
    ppr = paragraph._p.pPr
    assert ppr is not None
    paragraph_change = _change(ppr, "w:pPrChange")
    assert paragraph_change is not None, "无 pPr 的段落也必须补写 w:pPrChange"
    # OOXML 顺序：修订记录恒在 pPr / rPr 末尾
    assert ppr[-1].tag == qn("w:pPrChange")
    # 原本没有直接段落属性：旧状态是空 pPr，不能把新格式写进去，也不能伪造
    old_ppr = _old_state(paragraph_change, "w:pPr")
    assert len(old_ppr) == 0
    assert paragraph_change.get(qn("w:id"))
    assert paragraph_change.get(qn("w:author"))
    assert paragraph_change.get(qn("w:date"))

    run = paragraph.runs[0]
    rpr = run._element.find(qn("w:rPr"))
    assert rpr is not None
    run_change = _change(rpr, "w:rPrChange")
    assert run_change is not None, "无 rPr 的 run 也必须补写 w:rPrChange"
    assert rpr[-1].tag == qn("w:rPrChange")
    old_rpr = _old_state(run_change, "w:rPr")
    assert len(old_rpr) == 0


def test_old_snapshots_keep_original_values_not_new_format(tmp_path):
    """原有直接格式原值保留在旧快照里，新格式不得混入旧状态。"""
    source = tmp_path / "src.docx"
    tracked = tmp_path / "tracked.docx"

    document = Document()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("带旧直接格式的正文。")
    run.font.size = Pt(10)
    run.font.bold = True
    document.save(source)

    apply_format(str(source), BODY_SPEC, {0: "body"}, str(tracked), track=True)

    result = Document(tracked).paragraphs[0]
    old_ppr = _old_state(_change(result._p.pPr, "w:pPrChange"), "w:pPr")
    # 原对齐保留；新的两端对齐/缩进不得写进旧快照
    assert old_ppr.find(qn("w:jc")).get(qn("w:val")) == "center"
    assert old_ppr.find(qn("w:ind")) is None
    assert old_ppr.find(qn("w:spacing")) is None

    old_rpr = _old_state(
        _change(result.runs[0]._element.find(qn("w:rPr")), "w:rPrChange"),
        "w:rPr")
    assert old_rpr.find(qn("w:sz")).get(qn("w:val")) == "20"  # 原 10pt
    assert old_rpr.find(qn("w:sz")).get(qn("w:val")) != "32"  # 新 16pt
    assert old_rpr.find(qn("w:b")) is not None
    assert old_rpr.find(qn("w:rFonts")) is None


def test_label_prefix_split_keeps_revision_mapping_by_run_element(tmp_path):
    """label_prefix 拆出 run 后，旧快照仍按原 run 元素对应，不错位、不漏写。"""
    source = tmp_path / "src.docx"
    tracked = tmp_path / "tracked.docx"

    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("摘要：本文研究")
    first.font.size = Pt(10)
    second = paragraph.add_run("第二段文字")
    second.font.bold = True
    document.save(source)

    apply_format(str(source), LABEL_SPEC, {0: "body"}, str(tracked), track=True)

    paragraph = Document(tracked).paragraphs[0]
    runs = paragraph.runs
    # 前缀在第一个 run 内部，排版会把它拆成两个 run
    assert [run.text for run in runs] == ["摘要：", "本文研究", "第二段文字"]

    old_states = []
    for run in runs:
        rpr = run._element.find(qn("w:rPr"))
        assert rpr is not None
        change = _change(rpr, "w:rPrChange")
        assert change is not None, f"run {run.text!r} 缺少 w:rPrChange"
        assert rpr[-1].tag == qn("w:rPrChange")
        old_states.append(_old_state(change, "w:rPr"))

    # 第一个 run 是自己的旧属性（10pt、原本不加粗）
    assert old_states[0].find(qn("w:sz")).get(qn("w:val")) == "20"
    assert old_states[0].find(qn("w:b")) is None
    # 拆分出的新 run 没有旧状态：空 rPr，不能张冠李戴到第二个原 run
    assert len(old_states[1]) == 0
    # 第二个原 run 必须对应自己的旧快照（旧实现 zip 错位时会丢失）
    assert old_states[2].find(qn("w:b")) is not None


def test_clean_output_has_no_revision_markers(tmp_path):
    """track=False 干净稿不写任何修订记录。"""
    source = tmp_path / "src.docx"
    clean = tmp_path / "clean.docx"

    document = Document()
    document.add_paragraph("完全没有直接格式的正文。")
    document.save(source)

    apply_format(str(source), BODY_SPEC, {0: "body"}, str(clean), track=False)

    for paragraph in Document(clean).paragraphs:
        assert _change(paragraph._p.pPr, "w:pPrChange") is None
        for run in paragraph.runs:
            rpr = run._element.find(qn("w:rPr"))
            assert rpr is None or _change(rpr, "w:rPrChange") is None
